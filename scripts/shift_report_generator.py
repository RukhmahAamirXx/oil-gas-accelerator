#!/usr/bin/env python3
"""
shift_report_generator.py — KSEA Oil & Gas shift report (brief 5.2)
===================================================================
Generates an Excel shift report for a morning / afternoon / night shift:

  --shift morning    06:00 - 14:00
  --shift afternoon  14:00 - 22:00
  --shift night      22:00 - 06:00 (crosses midnight)

Queries Postgres for the shift window:
  * per-well anomaly summary (WARN / CRITICAL counts)
  * total maintenance cases opened during the shift
  * worst-performing well (most CRITICAL events)
  * anomaly frequency by hour (for the chart)

Builds an .xlsx with:
  * Summary sheet (shift metadata, totals, worst well)
  * Per-well anomaly table (WARN / CRITICAL / total), Excel formulas for totals
  * A bar chart of anomaly frequency by hour (openpyxl native chart)

Then:
  * Uploads the report to MinIO (boto3) — bucket kermit-shift-reports
  * Emails it to the shift supervisor via Mailhog (SMTP localhost:1025) with the
    .xlsx attached.

boto3 is used for MinIO (NOT the StreamSets S3 stage) so there is no checksum/
SSL-stage friction — a plain S3-compatible client to the MinIO endpoint.

Usage:
  python shift_report_generator.py --shift morning
  python shift_report_generator.py --shift night --date 2026-06-24
  python shift_report_generator.py --shift afternoon --no-upload --no-email
  python shift_report_generator.py --shift morning --local-only   # just write xlsx

Deps: psycopg2-binary, openpyxl, boto3
"""
import argparse
import datetime as dt
import io
import os
import smtplib
from email.message import EmailMessage
from pathlib import Path

# ---- Config ----
HERE = Path(__file__).resolve().parent
PG = {"host": "localhost", "port": 5434, "dbname": "oilgas", "user": "kermit", "password": "kermit"}

# MinIO (S3-compatible). Host-side endpoint; TLS with self-signed cert -> verify off.
MINIO = {
    "endpoint_url": "https://localhost:9000",
    "aws_access_key_id": "kermit",
    "aws_secret_access_key": "kermit123",
    "region_name": "us-east-1",
    "verify": False,   # self-signed cert on the MinIO container
}
MINIO_BUCKET = "kermit-shift-reports"

MAILHOG = {"host": "localhost", "port": 1025}
SUPERVISOR_EMAIL = "shift-supervisor@oilgas.local"
FROM_EMAIL = "shift-report@oilgas.local"

SHIFTS = {
    "morning":   (6, 14),
    "afternoon": (14, 22),
    "night":     (22, 6),   # crosses midnight
}


def shift_window(shift: str, base_date: dt.date):
    """Return (start_dt, end_dt) UTC for the shift on base_date."""
    start_h, end_h = SHIFTS[shift]
    start = dt.datetime.combine(base_date, dt.time(start_h, 0), tzinfo=dt.timezone.utc)
    if end_h > start_h:
        end = dt.datetime.combine(base_date, dt.time(end_h, 0), tzinfo=dt.timezone.utc)
    else:
        # night shift: ends next day at end_h
        end = dt.datetime.combine(base_date + dt.timedelta(days=1),
                                  dt.time(end_h, 0), tzinfo=dt.timezone.utc)
    return start, end


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------
def query_shift_data(cur, start, end):
    # per-well WARN/CRITICAL counts
    cur.execute(
        """
        SELECT well_id,
               COUNT(*) FILTER (WHERE severity = 'WARN')     AS warn_count,
               COUNT(*) FILTER (WHERE severity = 'CRITICAL') AS crit_count
        FROM anomaly_events
        WHERE event_ts >= %s AND event_ts < %s
        GROUP BY well_id
        ORDER BY crit_count DESC, warn_count DESC
        """,
        (start, end),
    )
    per_well = [(r[0], int(r[1]), int(r[2])) for r in cur.fetchall()]

    # maintenance cases opened during the shift
    cur.execute(
        "SELECT COUNT(*) FROM maintenance_cases WHERE recorded_at >= %s AND recorded_at < %s",
        (start, end),
    )
    cases_opened = int(cur.fetchone()[0])

    # anomaly frequency by hour
    cur.execute(
        """
        SELECT date_trunc('hour', event_ts) AS hr, COUNT(*)
        FROM anomaly_events
        WHERE event_ts >= %s AND event_ts < %s
        GROUP BY hr ORDER BY hr
        """,
        (start, end),
    )
    by_hour = [(r[0], int(r[1])) for r in cur.fetchall()]

    worst_well = per_well[0][0] if per_well else None
    return per_well, cases_opened, by_hour, worst_well


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def build_workbook(shift, base_date, start, end, per_well, cases_opened, by_hour, worst_well):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, Reference

    FONT = "Arial"
    hdr_fill = PatternFill("solid", start_color="1F4E78")
    hdr_font = Font(name=FONT, bold=True, color="FFFFFF")
    title_font = Font(name=FONT, bold=True, size=14)
    base_font = Font(name=FONT)
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()

    # ---- Summary sheet ----
    s = wb.active
    s.title = "Summary"
    s["A1"] = "Oil & Gas — Shift Report"
    s["A1"].font = title_font
    meta = [
        ("Shift", shift.capitalize()),
        ("Date", base_date.isoformat()),
        ("Window (UTC)", f"{start.strftime('%Y-%m-%d %H:%M')} → {end.strftime('%Y-%m-%d %H:%M')}"),
        ("Generated (UTC)", dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Maintenance cases opened", cases_opened),
        ("Worst-performing well", worst_well or "—"),
        ("Wells with anomalies", len(per_well)),
    ]
    r = 3
    for k, v in meta:
        s.cell(r, 1, k).font = Font(name=FONT, bold=True)
        s.cell(r, 2, v).font = base_font
        r += 1
    s.column_dimensions["A"].width = 26
    s.column_dimensions["B"].width = 40

    # ---- Per-well sheet ----
    w = wb.create_sheet("Per-Well Anomalies")
    headers = ["Well", "WARN", "CRITICAL", "Total"]
    for c, h in enumerate(headers, 1):
        cell = w.cell(1, c, h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    row = 2
    for well_id, warn, crit in per_well:
        w.cell(row, 1, well_id).font = base_font
        w.cell(row, 2, warn).font = base_font
        w.cell(row, 3, crit).font = base_font
        w.cell(row, 4, f"=B{row}+C{row}").font = base_font   # Excel formula for total
        for c in range(1, 5):
            w.cell(row, c).border = border
        row += 1
    # Totals row (formulas)
    if per_well:
        last = row - 1
        w.cell(row, 1, "TOTAL").font = Font(name=FONT, bold=True)
        w.cell(row, 2, f"=SUM(B2:B{last})").font = Font(name=FONT, bold=True)
        w.cell(row, 3, f"=SUM(C2:C{last})").font = Font(name=FONT, bold=True)
        w.cell(row, 4, f"=SUM(D2:D{last})").font = Font(name=FONT, bold=True)
        for c in range(1, 5):
            w.cell(row, c).border = border
    for col, width in zip("ABCD", (12, 10, 12, 10)):
        w.column_dimensions[col].width = width

    # ---- By-hour sheet + chart ----
    h = wb.create_sheet("Anomalies by Hour")
    h.cell(1, 1, "Hour (UTC)").font = hdr_font
    h.cell(1, 1).fill = hdr_fill
    h.cell(1, 2, "Anomaly Count").font = hdr_font
    h.cell(1, 2).fill = hdr_fill
    hr = 2
    for hour_dt, cnt in by_hour:
        h.cell(hr, 1, hour_dt.strftime("%H:%M")).font = base_font
        h.cell(hr, 2, cnt).font = base_font
        hr += 1
    h.column_dimensions["A"].width = 12
    h.column_dimensions["B"].width = 16

    if by_hour:
        chart = BarChart()
        chart.title = f"Anomaly Frequency by Hour — {shift.capitalize()} Shift"
        chart.x_axis.title = "Hour (UTC)"
        chart.y_axis.title = "Anomalies"
        chart.height = 9
        chart.width = 18
        data = Reference(h, min_col=2, min_row=1, max_row=hr - 1)
        cats = Reference(h, min_col=1, min_row=2, max_row=hr - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend = None
        h.add_chart(chart, "D2")

    return wb


# ---------------------------------------------------------------------------
# Upload + email
# ---------------------------------------------------------------------------
def upload_to_minio(data_bytes, key):
    import boto3
    from botocore.config import Config
    s3 = boto3.client(
        "s3",
        endpoint_url=MINIO["endpoint_url"],
        aws_access_key_id=MINIO["aws_access_key_id"],
        aws_secret_access_key=MINIO["aws_secret_access_key"],
        region_name=MINIO["region_name"],
        verify=MINIO["verify"],
        config=Config(s3={"addressing_style": "path"}),   # MinIO needs path-style
    )
    # ensure bucket
    try:
        s3.head_bucket(Bucket=MINIO_BUCKET)
    except Exception:
        s3.create_bucket(Bucket=MINIO_BUCKET)
    s3.put_object(Bucket=MINIO_BUCKET, Key=key, Body=data_bytes,
                  ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return f"s3://{MINIO_BUCKET}/{key}"


def email_report(data_bytes, filename, shift, base_date, cases_opened, worst_well):
    msg = EmailMessage()
    msg["Subject"] = f"Shift Report — {shift.capitalize()} — {base_date.isoformat()}"
    msg["From"] = FROM_EMAIL
    msg["To"] = SUPERVISOR_EMAIL
    msg.set_content(
        f"Attached is the {shift} shift report for {base_date.isoformat()}.\n\n"
        f"Maintenance cases opened this shift: {cases_opened}\n"
        f"Worst-performing well: {worst_well or '—'}\n"
    )
    msg.add_attachment(
        data_bytes, maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
    )
    with smtplib.SMTP(MAILHOG["host"], MAILHOG["port"], timeout=10) as smtp:
        smtp.send_message(msg)


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Shift report generator (brief 5.2)")
    ap.add_argument("--shift", required=True, choices=list(SHIFTS.keys()))
    ap.add_argument("--date", help="base date YYYY-MM-DD (default: today UTC)")
    ap.add_argument("--no-upload", action="store_true", help="skip MinIO upload")
    ap.add_argument("--no-email", action="store_true", help="skip Mailhog email")
    ap.add_argument("--local-only", action="store_true", help="write xlsx locally, no upload/email")
    ap.add_argument("--outdir", default=str(HERE / "reports"))
    args = ap.parse_args()

    base_date = (dt.date.fromisoformat(args.date) if args.date
                 else dt.datetime.now(dt.timezone.utc).date())
    start, end = shift_window(args.shift, base_date)

    import psycopg2
    conn = psycopg2.connect(**PG)
    cur = conn.cursor()
    per_well, cases_opened, by_hour, worst_well = query_shift_data(cur, start, end)
    cur.close(); conn.close()

    print(f"Shift {args.shift} {base_date}: {len(per_well)} wells with anomalies, "
          f"{cases_opened} cases, worst={worst_well}")

    wb = build_workbook(args.shift, base_date, start, end,
                        per_well, cases_opened, by_hour, worst_well)
    buf = io.BytesIO()
    wb.save(buf)
    data_bytes = buf.getvalue()

    filename = f"shift_report_{args.shift}_{base_date.isoformat()}.xlsx"

    # always write a local copy
    outdir = Path(args.outdir); outdir.mkdir(parents=True, exist_ok=True)
    local_path = outdir / filename
    local_path.write_bytes(data_bytes)
    print(f"  wrote {local_path}")

    if args.local_only:
        return

    if not args.no_upload:
        try:
            uri = upload_to_minio(data_bytes, filename)
            print(f"  uploaded {uri}")
        except Exception as e:
            print(f"  MinIO upload failed: {e}")

    if not args.no_email:
        try:
            email_report(data_bytes, filename, args.shift, base_date, cases_opened, worst_well)
            print(f"  emailed to {SUPERVISOR_EMAIL} via Mailhog")
        except Exception as e:
            print(f"  Mailhog email failed: {e}")


if __name__ == "__main__":
    main()
